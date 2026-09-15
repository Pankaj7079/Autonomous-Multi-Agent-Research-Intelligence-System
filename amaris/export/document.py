"""Renders a finished report as a .docx.

Walks markdown-it's token stream rather than re-parsing the markdown, because markdown-it
is already core and already renders the same report in the UI — a second dialect here would
drift from what the user saw on screen.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from typing import TYPE_CHECKING, Any

from amaris.observability.logging import logger

if TYPE_CHECKING:
    from amaris.api.schemas import ResearchResult

MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# long enough to tell two downloads apart, short enough to stay a filename
SLUG_CHARS = 40
# word's built-in heading styles stop being visually distinct below this
MAX_HEADING_LEVEL = 4
CODE_FONT = "Consolas"
CODE_POINTS = 9
META_POINTS = 9

_SLUG = re.compile(r"[^a-z0-9]+")


class ExportUnavailable(RuntimeError):
    """python-docx is not installed — the export extra was not synced."""


@dataclass(frozen=True)
class Block:
    """One renderable piece of the report, flattened out of the token stream."""

    kind: str
    text: str
    level: int = 0


def _inline_text(token: Any) -> str:
    """One inline token as plain text, with link targets kept.

    Emphasis is flattened: a docx run per styled span is buildable but buys little in a
    report that is almost entirely plain prose. Links keep their url inline because a
    printed document gives the reader no way to hover one.
    """
    parts: list[str] = []
    href = ""
    for child in token.children or []:
        if child.type in ("text", "code_inline"):
            parts.append(child.content)
        elif child.type == "softbreak":
            parts.append(" ")
        elif child.type == "hardbreak":
            parts.append("\n")
        elif child.type == "link_open":
            href = child.attrGet("href") or ""
        elif child.type == "link_close":
            # skip when the text already is the url, or every reference reads "url (url)"
            if href and href not in "".join(parts):
                parts.append(f" ({href})")
            href = ""
    return "".join(parts).strip()


def to_blocks(markdown_text: str) -> list[Block]:
    """Report markdown flattened to blocks in document order."""
    from markdown_it import MarkdownIt

    tokens = MarkdownIt("commonmark").enable("table").parse(markdown_text or "")
    blocks: list[Block] = []
    heading = ""
    lists: list[str] = []
    quoted = 0
    # table cells arrive as ordinary inline tokens, so a row is collected and joined
    row: list[str] | None = None

    for token in tokens:
        kind = token.type
        if kind == "heading_open":
            heading = token.tag
        elif kind == "heading_close":
            heading = ""
        elif kind == "bullet_list_open":
            lists.append("bullet")
        elif kind == "ordered_list_open":
            lists.append("ordered")
        elif kind in ("bullet_list_close", "ordered_list_close"):
            if lists:
                lists.pop()
        elif kind == "blockquote_open":
            quoted += 1
        elif kind == "blockquote_close":
            quoted = max(0, quoted - 1)
        elif kind == "tr_open":
            row = []
        elif kind == "tr_close":
            if row:
                blocks.append(Block("body", " | ".join(row)))
            row = None
        elif kind in ("fence", "code_block"):
            blocks.append(Block("code", token.content.rstrip("\n")))
        elif kind == "hr":
            blocks.append(Block("rule", ""))
        elif kind == "inline":
            text = _inline_text(token)
            if not text:
                continue
            if row is not None:
                row.append(text)
            elif heading:
                blocks.append(Block("heading", text, int(heading[1:])))
            elif lists:
                blocks.append(Block(lists[-1], text, len(lists)))
            elif quoted:
                blocks.append(Block("quote", text))
            else:
                blocks.append(Block("body", text))
    return blocks


def _meta_line(result: ResearchResult | None, session_id: str, elapsed: float) -> str:
    """The run's own numbers, so the document says how it was produced."""
    parts = [datetime.now(UTC).strftime("%d %B %Y")]
    trace = getattr(result, "trace", None)
    if trace is not None:
        depth = str(trace.triage.get("depth", "")) if trace.triage else ""
        if depth:
            parts.append(f"{depth} research")
        parts.append(f"quality {trace.quality_score:.2f}")
        parts.append(f"{trace.source_count} sources")
    if elapsed:
        parts.append(f"{elapsed:.0f}s")
    if session_id:
        parts.append(f"run {session_id[:8]}")
    return "  ·  ".join(parts)


def build_markdown_docx(markdown_text: str, title: str, subtitle: str = "") -> bytes:
    """Any markdown as .docx bytes. Raises ExportUnavailable without the export extra."""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise ExportUnavailable("install the export extra: uv sync --extra export") from exc

    document = Document()
    document.add_heading(title or "AMARIS report", level=0)

    if subtitle:
        meta = document.add_paragraph()
        run = meta.add_run(subtitle)
        run.italic = True
        run.font.size = Pt(META_POINTS)

    for block in to_blocks(markdown_text):
        if block.kind == "heading":
            document.add_heading(block.text, level=min(block.level, MAX_HEADING_LEVEL))
        elif block.kind == "bullet":
            document.add_paragraph(block.text, style="List Bullet")
        elif block.kind == "ordered":
            document.add_paragraph(block.text, style="List Number")
        elif block.kind == "quote":
            document.add_paragraph(block.text, style="Intense Quote")
        elif block.kind == "code":
            code = document.add_paragraph().add_run(block.text)
            code.font.name = CODE_FONT
            code.font.size = Pt(CODE_POINTS)
        elif block.kind == "rule":
            document.add_paragraph()
        else:
            document.add_paragraph(block.text)

    buffer = BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    logger.bind(bytes=len(data), chars=len(markdown_text)).info("export.docx_built")
    return data


def build_docx(
    result: ResearchResult | None,
    query: str,
    *,
    session_id: str = "",
    elapsed: float = 0.0,
) -> bytes:
    """One finished turn as .docx bytes, titled by its question."""
    return build_markdown_docx(
        result.report if result is not None else "",
        query.strip(),
        _meta_line(result, session_id, elapsed),
    )


def docx_filename(query: str) -> str:
    """A filename the user can recognise in a downloads folder six files later."""
    slug = _SLUG.sub("-", query.lower()).strip("-")[:SLUG_CHARS].strip("-")
    return f"amaris_{slug or 'report'}_{datetime.now(UTC):%Y%m%d}.docx"
